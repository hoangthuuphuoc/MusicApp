from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"D:\MusicApp2\MusicApp_Huong_Dan_Luong_Home.docx"
NAVY = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side_name in ("top", "left", "bottom", "right"):
        side = borders.find(qn(f"w:{side_name}"))
        if side is None:
            side = OxmlElement(f"w:{side_name}")
            borders.append(side)
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), "4")
        side.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="Arial", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=7, line_spacing=1.18):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_font(run)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=15 if level == 1 else 12, bold=True)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, space_after=3)
    run = p.add_run(text)
    set_font(run)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, space_after=4)
    run = p.add_run(text)
    set_font(run)
    return p


def add_flow(doc, lines):
    table = doc.add_table(rows=len(lines), cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.55)
    for index, line in enumerate(lines):
        cell = table.cell(index, 0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, LIGHT_BLUE if index % 2 == 0 else LIGHT_GRAY)
        set_cell_border(cell)
        set_cell_margin(cell, 90, 140, 90, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=0, line_spacing=1.0)
        run = p.add_run(line)
        set_font(run, name="Consolas", size=10, bold=True, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].width = Inches(widths[i])
        header_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(header_cells[i], NAVY)
        set_cell_border(header_cells[i])
        set_cell_margin(header_cells[i])
        p = header_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=0, line_spacing=1.0)
        run = p.add_run(header)
        set_font(run, size=10, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cells[i], "FFFFFF" if row_index % 2 == 0 else LIGHT_BLUE)
            set_cell_border(cells[i])
            set_cell_margin(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, space_after=0, line_spacing=1.08)
            run = p.add_run(value)
            set_font(run, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_font(run, name="Consolas", size=9, color="17365D")
    return p


def make_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Huong Dan Luong Chay MusicApp Tu Man Hinh Home")
    set_font(title_run, size=24, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("Tai lieu ky thuat ve phan Home Search Detail va phat nhac")
    set_font(subtitle_run, size=11, color="4F4F4F")

    add_body(
        doc,
        "Tài liệu này giải thích cách ứng dụng MusicApp hoạt động kể từ khi người dùng vào màn hình Home. Trọng tâm là luồng tải nhạc từ thiết bị, tìm kiếm, mở chi tiết bài hát và điều khiển phát nhạc. Phần Splash và Onboarding không nằm trong phạm vi tài liệu.",
    )

    add_heading(doc, "Pham vi va dieu kien bat dau")
    add_body(doc, "Luồng được mô tả bắt đầu khi MainActivity đã mở và ViewPager hiển thị HomeFragment. Trước khi tải danh sách nhạc, ứng dụng cần được người dùng cho phép đọc audio. Trên Android 13 trở lên, quyền chính là READ MEDIA AUDIO; trên Android 12 trở xuống, ứng dụng dùng READ EXTERNAL STORAGE.")
    add_table(
        doc,
        ["Thanh phan", "Vai tro tu Home tro di"],
        [
            ["MainActivity", "Chứa ViewPager và BottomNavigation; trang 0 là Home, trang 1 là Search."],
            ["HomeFragment", "Hiển thị album nổi bật, danh sách bài hát và mini player."],
            ["SearchFragment", "Lọc danh sách nhạc theo nội dung người dùng nhập."],
            ["DetailActivity", "Hiển thị thông tin bài hát, tiến trình và các nút điều khiển phát."],
            ["MusicPlaybackService", "Chạy nền dạng foreground service, điều khiển MediaPlayer và notification."],
        ],
        [1.75, 4.8],
    )

    add_heading(doc, "Kien truc chinh")
    add_body(doc, "Phần chức năng sau Home dùng mô hình tách lớp: giao diện gửi sự kiện đến ViewModel; ViewModel gọi repository; repository đọc MediaStore hoặc gửi action tới service. Trạng thái phát nhạc được phát lại qua StateFlow để Home và Detail cùng cập nhật giao diện.")
    add_flow(doc, [
        "HomeFragment -> HomeViewModel -> AudioRepository -> AudioLocalDataSource -> MediaStore",
        "Nhan bai hat -> HomeViewModel -> MusicPlaybackRepository -> MusicPlaybackService -> MediaPlayer",
        "MusicPlaybackService -> playbackState StateFlow -> HomeViewModel va DetailViewModel -> giao dien",
    ])

    add_heading(doc, "Luồng tai danh sach nhac o Home")
    add_number(doc, "MainActivity tạo MainPagerAdapter. Adapter trả về HomeFragment cho trang đầu tiên của ViewPager.")
    add_number(doc, "Trong HomeFragment.onViewCreated, RecyclerView được cấu hình cho danh sách album ngang và danh sách bài hát dọc. Sau đó fragment gửi HomeUiEvent LoadAudio.")
    add_number(doc, "HomeViewModel nhận sự kiện và gọi AudioRepository.getAudios trong viewModelScope. Repository chuyển việc đọc dữ liệu cho AudioLocalDataSource.")
    add_number(doc, "AudioLocalDataSource truy vấn MediaStore Audio với điều kiện IS MUSIC khác 0 và thời lượng lớn hơn 0. Mỗi dòng dữ liệu được chuyển thành model Audio gồm id, tiêu đề, nghệ sĩ, URI file, ảnh album và thời lượng.")
    add_number(doc, "HomeViewModel cập nhật HomeUiState. HomeFragment đang collect state trong repeatOnLifecycle nên submit danh sách mới cho HomeFeaturedAdapter và HomeSongAdapter.")
    add_code(doc, "HomeFragment -> HomeUiEvent.LoadAudio -> HomeViewModel.loadAudio -> AudioRepository.getAudios -> MediaStore")

    add_heading(doc, "Trang thai giao dien Home")
    add_table(
        doc,
        ["State", "Noi dung", "Noi su dung"],
        [
            ["audios", "Toàn bộ audio đọc được từ thiết bị.", "Danh sách Trend Streams."],
            ["featuredAudios", "Tối đa 10 bài hát đầu tiên trong danh sách.", "Danh sách album ngang."],
            ["playbackState", "Bài hát hiện tại, đang phát hay không, vị trí và thời lượng.", "Mini player và dữ liệu đồng bộ sang Detail."],
            ["isLoading va error", "Trạng thái đang tải hoặc lỗi khi truy vấn MediaStore.", "Sẵn sàng để giao diện hiển thị phản hồi phù hợp."],
        ],
        [1.35, 3.2, 2.0],
    )

    add_heading(doc, "Luồng tim kiem")
    add_body(doc, "Khi chạm ô tìm kiếm trên Home, HomeViewModel phát HomeUiEffect OpenSearch. HomeFragment gọi MainActivity.openSearch, hàm này chuyển BottomNavigation và ViewPager sang trang Search.")
    add_body(doc, "SearchFragment cũng tải danh sách audio khi khởi tạo. Mỗi lần nội dung EditText thay đổi, fragment gửi SearchUiEvent Search. SearchViewModel lọc danh sách hiện có và trả kết quả qua SearchUiState.resultAudios. Nhấn Cancel sẽ xóa text và gửi ClearSearch để trở về danh sách phù hợp.")
    add_code(doc, "Home search box -> HomeUiEffect.OpenSearch -> MainActivity.openSearch -> SearchFragment")

    add_heading(doc, "Luồng chon bai hat va phat nhac")
    add_number(doc, "Người dùng chạm một bài trong Home hoặc Search. Adapter gọi callback và ViewModel nhận ClickAudio cùng Audio đã chọn.")
    add_number(doc, "ViewModel gọi MusicPlaybackRepository.play(audioId). Repository tạo Intent action play và khởi động MusicPlaybackService bằng startForegroundService.")
    add_number(doc, "Service tải danh sách audio nếu cần, xác định vị trí của audioId, tạo MediaPlayer, đặt URI bài hát và gọi prepareAsync.")
    add_number(doc, "Khi MediaPlayer chuẩn bị xong, service bắt đầu phát, cập nhật PlaybackState isPlaying, thời lượng và vị trí. Service cũng tạo notification có Previous, Play Pause và Next.")
    add_number(doc, "HomeViewModel và DetailViewModel collect MusicPlaybackService.playbackState. Giao diện đổi tên bài hát, nghệ sĩ, ảnh album, tiến trình và icon play pause theo trạng thái mới.")
    add_flow(doc, [
        "ClickAudio -> MusicPlaybackRepository.play audioId",
        "MusicPlaybackService PLAY -> MediaPlayer prepareAsync -> start",
        "playbackState cap nhat moi 500 ms -> Mini player va Detail cap nhat",
    ])

    add_heading(doc, "Mini player va man hinh Detail")
    add_body(doc, "Mini player ở cuối Home hiển thị bài hát đang phát. Nút của nó gửi HomeUiEvent ClickPlayPause; ViewModel chuyển tiếp thành action play pause đến service. Khi đang phát, giao diện dùng icon pause; khi tạm dừng, giao diện dùng icon play.")
    add_body(doc, "Sau khi chọn bài hát, Home hoặc Search phát effect OpenDetail. DetailActivity nhận audio id từ Intent, gửi DetailUiEvent LoadAudio và hiển thị thông tin bài hát. Các nút Back, Play Pause, Next, Previous và SeekBar đều được ánh xạ thành DetailUiEvent rồi gửi đến MusicPlaybackRepository.")
    add_table(
        doc,
        ["Thao tac", "Action gui den service", "Ket qua"],
        [
            ["Play Pause", "play pause", "Tạm dừng hoặc tiếp tục MediaPlayer và cập nhật notification."],
            ["Next", "next", "Tăng index bài hát, quay về đầu danh sách khi đã hết."],
            ["Previous", "previous", "Giảm index bài hát, quay về cuối danh sách khi đang ở đầu."],
            ["Keo SeekBar", "seek kèm vị trí mili giây", "Service giới hạn vị trí hợp lệ rồi MediaPlayer.seekTo."],
            ["Notification", "previous play pause next", "PendingIntent gọi lại service khi người dùng thao tác ngoài ứng dụng."],
        ],
        [1.35, 2.25, 2.95],
    )

    add_heading(doc, "Cac file nen doc khi bao tri")
    add_table(
        doc,
        ["File", "Ly do doc"],
        [
            ["MainActivity.kt", "Điểm vào phần Home và Search; quản lý ViewPager cùng BottomNavigation."],
            ["ui fragment home HomeFragment.kt", "Liên kết layout Home, adapter, state và effect."],
            ["ui home HomeViewModel.kt", "Tải danh sách audio, xử lý thao tác Home và quan sát playbackState."],
            ["data local AudioLocalDataSource.kt", "Nguồn dữ liệu nhạc cục bộ từ MediaStore."],
            ["data repository MusicPlaybackRepository.kt", "Cổng gửi action từ ViewModel đến service."],
            ["service MusicPlaybackService.kt", "Logic MediaPlayer, foreground service, tiến trình và notification."],
            ["ui detail DetailActivity.kt va DetailViewModel.kt", "Giao diện điều khiển và đồng bộ trạng thái bài hát."],
        ],
        [2.65, 3.9],
    )

    add_heading(doc, "Kiem tra nhanh khi chay thu")
    add_bullet(doc, "Cấp quyền đọc nhạc và quyền notification khi hệ thống yêu cầu.")
    add_bullet(doc, "Mở Home và xác nhận danh sách nhạc từ thiết bị xuất hiện.")
    add_bullet(doc, "Chọn một bài hát, kiểm tra Detail mở ra và notification xuất hiện.")
    add_bullet(doc, "Thử Play Pause, Next, Previous và kéo thanh tiến trình.")
    add_bullet(doc, "Quay về Home, xác nhận mini player phản ánh đúng bài hát và trạng thái hiện tại.")
    add_bullet(doc, "Mở Search, nhập một phần tên bài hát hoặc nghệ sĩ, sau đó chọn kết quả để kiểm tra cùng luồng phát nhạc.")

    doc.save(OUTPUT)


if __name__ == "__main__":
    make_document()
